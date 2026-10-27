#!/usr/bin/env python3
"""
美甲图片匹配器 Web 应用
"""

import os
import sys
import sqlite3
import tempfile
import webbrowser
import threading
import time
from pathlib import Path
from flask import Flask, request, jsonify, render_template, send_from_directory, send_file
from werkzeug.utils import secure_filename
import numpy as np
from PIL import Image
import io

def get_base_path():
    """获取应用基础路径，兼容PyInstaller打包"""
    if getattr(sys, 'frozen', False):
        return Path(sys._MEIPASS)
    else:
        return Path(__file__).parent

BASE_PATH = get_base_path()
PROJECT_ROOT = BASE_PATH.parent if not getattr(sys, 'frozen', False) else BASE_PATH
DB_PATH = BASE_PATH / "data" / "货盘分类.db"
IMAGES_DIR = BASE_PATH / "data" / "货盘分类图片"
MODEL_PATH = BASE_PATH / "models" / "sentence-transformers" / "clip-ViT-B-32"

app = Flask(__name__,
            template_folder='templates',
            static_folder='static',
            static_url_path='/static')

app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024
app.config['UPLOAD_FOLDER'] = tempfile.mkdtemp()

ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'webp', 'gif'}

model = None

def check_model():
    """检查模型是否存在"""
    model_path = MODEL_PATH
    if not model_path.exists():
        print("=" * 50)
        print("错误: AI模型未找到！")
        print("=" * 50)
        print(f"期望路径: {model_path}")
        print()
        print("请先下载模型:")
        print("  python download_model.py")
        print()
        print("或手动下载:")
        print("  https://modelscope.cn/models/sentence-transformers/clip-ViT-B-32/files")
        print("=" * 50)
        sys.exit(1)

def get_model():
    global model
    if model is None:
        print("加载模型...")
        from sentence_transformers import SentenceTransformer
        model = SentenceTransformer(str(MODEL_PATH))
        print("模型加载完成")
    return model

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def search_similar(image_path, top_k=10):
    matcher = get_model()
    
    image = Image.open(image_path).convert('RGB')
    query_features = matcher.encode(image).astype(np.float32)
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute('''
        SELECT f.item_id, f.feature_vector, f.image_path, i.item_code, i.price, c.name as category
        FROM image_features f
        JOIN items i ON f.item_id = i.id
        JOIN categories c ON i.category_id = c.id
    ''')
    
    results = []
    for item_id, feat_blob, img_path, item_code, price, category in cursor.fetchall():
        feat = np.frombuffer(feat_blob, dtype=np.float32)
        similarity = np.dot(query_features, feat) / (
            np.linalg.norm(query_features) * np.linalg.norm(feat)
        )
        results.append({
            'item_id': item_id,
            'item_code': item_code,
            'price': price,
            'category': category,
            'image_path': img_path,
            'similarity': float(similarity)
        })
    
    conn.close()
    results.sort(key=lambda x: x['similarity'], reverse=True)
    return results[:top_k]

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/search', methods=['POST'])
def api_search():
    if 'image' not in request.files:
        return jsonify({'error': '没有上传图片'}), 400
    
    file = request.files['image']
    if file.filename == '':
        return jsonify({'error': '没有选择文件'}), 400
    
    if not allowed_file(file.filename):
        return jsonify({'error': '不支持的文件格式'}), 400
    
    top_k = request.form.get('top_k', 10, type=int)
    
    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)
    
    try:
        results = search_similar(filepath, top_k)
        return jsonify({'results': results})
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        if os.path.exists(filepath):
            os.remove(filepath)

@app.route('/images/<filename>')
def serve_image(filename):
    return send_from_directory(IMAGES_DIR, filename)

@app.route('/admin')
def admin():
    return render_template('admin.html')

@app.route('/api/categories', methods=['GET'])
def api_categories():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute('''
        SELECT c.id, c.name, c.code, COUNT(i.id) as item_count
        FROM categories c
        LEFT JOIN items i ON i.category_id = c.id
        GROUP BY c.id
        ORDER BY c.id
    ''')
    
    categories = [{
        'id': row[0],
        'name': row[1],
        'code': row[2],
        'item_count': row[3]
    } for row in cursor.fetchall()]
    
    conn.close()
    return jsonify({'categories': categories})

@app.route('/api/categories', methods=['POST'])
def api_category_create():
    data = request.get_json()
    
    name = data.get('name')
    code = data.get('code')
    
    if not name:
        return jsonify({'error': '分类名称为必填项'}), 400
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute('SELECT id FROM categories WHERE name = ?', (name,))
    if cursor.fetchone():
        conn.close()
        return jsonify({'error': '分类名称已存在'}), 400
    
    cursor.execute('INSERT INTO categories (name, code) VALUES (?, ?)', (name, code))
    category_id = cursor.lastrowid
    conn.commit()
    conn.close()
    
    return jsonify({'id': category_id, 'message': '分类添加成功'})

@app.route('/api/categories/<int:category_id>', methods=['PUT'])
def api_category_update(category_id):
    data = request.get_json()
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute('SELECT id FROM categories WHERE id = ?', (category_id,))
    if not cursor.fetchone():
        conn.close()
        return jsonify({'error': '分类不存在'}), 404
    
    updates = []
    params = []
    
    if 'name' in data:
        cursor.execute('SELECT id FROM categories WHERE name = ? AND id != ?', (data['name'], category_id))
        if cursor.fetchone():
            conn.close()
            return jsonify({'error': '分类名称已存在'}), 400
        updates.append('name = ?')
        params.append(data['name'])
    
    if 'code' in data:
        updates.append('code = ?')
        params.append(data['code'])
    
    if not updates:
        conn.close()
        return jsonify({'error': '没有要更新的内容'}), 400
    
    params.append(category_id)
    cursor.execute(f'UPDATE categories SET {", ".join(updates)} WHERE id = ?', params)
    conn.commit()
    conn.close()
    
    return jsonify({'message': '分类更新成功'})

@app.route('/api/categories/<int:category_id>', methods=['DELETE'])
def api_category_delete(category_id):
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute('SELECT id FROM categories WHERE id = ?', (int(category_id),))
    if not cursor.fetchone():
        conn.close()
        return jsonify({'error': '分类不存在'}), 404
    
    cursor.execute('SELECT COUNT(*) FROM items WHERE category_id = ?', (int(category_id),))
    item_count = cursor.fetchone()[0]
    
    if item_count > 0:
        conn.close()
        return jsonify({'error': f'该分类下有{item_count}个货品，无法删除'}), 400
    
    cursor.execute('DELETE FROM categories WHERE id = ?', (int(category_id),))
    conn.commit()
    conn.close()
    
    return jsonify({'message': '分类删除成功'})

@app.route('/api/items', methods=['GET'])
def api_items_list():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 20, type=int)
    category_id = request.args.get('category_id', type=int)
    search = request.args.get('search', '')
    
    offset = (page - 1) * per_page
    
    query = '''
        SELECT i.id, i.item_code, i.price, i.category_id, i.image_path, c.name as category
        FROM items i
        JOIN categories c ON i.category_id = c.id
    '''
    params = []
    
    if category_id:
        query += ' WHERE i.category_id = ?'
        params.append(category_id)
    
    if search:
        if category_id:
            query += ' AND i.item_code LIKE ?'
        else:
            query += ' WHERE i.item_code LIKE ?'
        params.append(f'%{search}%')
    
    query += f' ORDER BY i.id LIMIT ? OFFSET ?'
    params.extend([per_page, offset])
    
    cursor.execute(query, params)
    items = [{
        'id': row[0],
        'item_code': row[1],
        'price': row[2],
        'category_id': row[3],
        'image_path': row[4],
        'category': row[5]
    } for row in cursor.fetchall()]
    
    cursor.execute('SELECT COUNT(*) FROM items')
    total = cursor.fetchone()[0]
    
    conn.close()
    return jsonify({'items': items, 'total': total, 'page': page, 'per_page': per_page})

@app.route('/api/items/<int:item_id>', methods=['GET'])
def api_item_detail(item_id):
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute('''
        SELECT i.id, i.item_code, i.price, i.category_id, i.image_path, c.name as category
        FROM items i
        JOIN categories c ON i.category_id = c.id
        WHERE i.id = ?
    ''', (item_id,))
    row = cursor.fetchone()
    conn.close()
    
    if not row:
        return jsonify({'error': '货品不存在'}), 404
    
    return jsonify({
        'id': row[0],
        'item_code': row[1],
        'price': row[2],
        'category_id': row[3],
        'image_path': row[4],
        'category': row[5]
    })

@app.route('/api/items', methods=['POST'])
def api_item_create():
    data = request.get_json()
    
    item_code = data.get('item_code')
    price = data.get('price')
    category_id = data.get('category_id')
    
    if not item_code or not category_id:
        return jsonify({'error': '货号和分类为必填项'}), 400
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute('''
        INSERT INTO items (item_code, price, category_id)
        VALUES (?, ?, ?)
    ''', (item_code, price, category_id))
    
    item_id = cursor.lastrowid
    conn.commit()
    conn.close()
    
    return jsonify({'id': item_id, 'message': '货品添加成功'})

@app.route('/api/items/<int:item_id>', methods=['PUT'])
def api_item_update(item_id):
    data = request.get_json()
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute('SELECT id FROM items WHERE id = ?', (item_id,))
    if not cursor.fetchone():
        conn.close()
        return jsonify({'error': '货品不存在'}), 404
    
    updates = []
    params = []
    
    if 'item_code' in data:
        updates.append('item_code = ?')
        params.append(data['item_code'])
    
    if 'price' in data:
        updates.append('price = ?')
        params.append(data['price'])
    
    if 'category_id' in data:
        updates.append('category_id = ?')
        params.append(data['category_id'])
    
    if not updates:
        conn.close()
        return jsonify({'error': '没有要更新的内容'}), 400
    
    params.append(item_id)
    cursor.execute(f'UPDATE items SET {", ".join(updates)} WHERE id = ?', params)
    conn.commit()
    conn.close()
    
    return jsonify({'message': '货品更新成功'})

@app.route('/api/items/<int:item_id>', methods=['DELETE'])
def api_item_delete(item_id):
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute('SELECT id FROM items WHERE id = ?', (item_id,))
    if not cursor.fetchone():
        conn.close()
        return jsonify({'error': '货品不存在'}), 404
    
    cursor.execute('DELETE FROM image_features WHERE item_id = ?', (item_id,))
    cursor.execute('DELETE FROM items WHERE id = ?', (item_id,))
    conn.commit()
    conn.close()
    
    return jsonify({'message': '货品删除成功'})

@app.route('/api/items/<int:item_id>/image', methods=['POST'])
def api_item_upload_image(item_id):
    if 'image' not in request.files:
        return jsonify({'error': '没有上传图片'}), 400
    
    file = request.files['image']
    if file.filename == '':
        return jsonify({'error': '没有选择文件'}), 400
    
    if not allowed_file(file.filename):
        return jsonify({'error': '不支持的文件格式'}), 400
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute('SELECT id FROM items WHERE id = ?', (item_id,))
    if not cursor.fetchone():
        conn.close()
        return jsonify({'error': '货品不存在'}), 404
    
    ext = file.filename.rsplit('.', 1)[1].lower()
    new_filename = f'item_{item_id}.{ext}'
    filepath = os.path.join(IMAGES_DIR, new_filename)
    file.save(filepath)
    
    relative_path = f'货盘分类图片/{new_filename}'
    cursor.execute('UPDATE items SET image_path = ? WHERE id = ?', (relative_path, item_id))
    
    matcher = get_model()
    image = Image.open(filepath).convert('RGB')
    features = matcher.encode(image).astype(np.float32)
    
    cursor.execute('DELETE FROM image_features WHERE item_id = ?', (item_id,))
    cursor.execute('''
        INSERT INTO image_features (item_id, image_path, feature_vector)
        VALUES (?, ?, ?)
    ''', (item_id, relative_path, features.tobytes()))
    
    conn.commit()
    conn.close()
    
    return jsonify({'message': '图片上传成功', 'image_path': relative_path})

@app.route('/api/export', methods=['GET'])
def api_export():
    """导出货品数据为Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.drawing.image import Image as XLImage
    except ImportError:
        return jsonify({'error': '需要安装openpyxl: pip install openpyxl'}), 500
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    category_id = request.args.get('category_id', type=int)
    search = request.args.get('search', '')
    include_images = request.args.get('images', 'true').lower() == 'true'
    
    query = '''
        SELECT i.id, i.item_code, i.price, c.name as category, c.code as category_code, i.image_path
        FROM items i
        JOIN categories c ON i.category_id = c.id
    '''
    params = []
    
    if category_id:
        query += ' WHERE i.category_id = ?'
        params.append(category_id)
    
    if search:
        if category_id:
            query += ' AND i.item_code LIKE ?'
        else:
            query += ' WHERE i.item_code LIKE ?'
        params.append(f'%{search}%')
    
    query += ' ORDER BY c.id, i.id'
    
    cursor.execute(query, params)
    items = cursor.fetchall()
    conn.close()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '货品列表'
    
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color='FBCFE8', end_color='FBCFE8', fill_type='solid')
    
    if include_images:
        headers = ['图片', 'ID', '货号', '价格', '分类', '分类编号']
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 8
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 10
    else:
        headers = ['ID', '货号', '价格', '分类', '分类编号', '图片路径']
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 10
        ws.column_dimensions['F'].width = 30
    
    ws.append(headers)
    for col_idx, cell in enumerate(ws[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    img_size = 100
    
    for row_idx, item in enumerate(items, 2):
        item_id, item_code, price, category, category_code, image_path = item
        
        if include_images:
            ws.row_dimensions[row_idx].height = img_size * 0.75
            ws.cell(row=row_idx, column=2, value=item_id)
            ws.cell(row=row_idx, column=3, value=item_code)
            ws.cell(row=row_idx, column=4, value=price)
            ws.cell(row=row_idx, column=5, value=category)
            ws.cell(row=row_idx, column=6, value=category_code)
            
            if image_path:
                img_filename = image_path.split('/')[-1]
                img_full_path = IMAGES_DIR / img_filename
                
                if img_full_path.exists():
                    try:
                        img = XLImage(str(img_full_path))
                        img.width = img_size
                        img.height = img_size
                        ws.add_image(img, f'A{row_idx}')
                    except Exception:
                        ws.cell(row=row_idx, column=1, value='×')
                else:
                    ws.cell(row=row_idx, column=1, value='×')
            else:
                ws.cell(row=row_idx, column=1, value='×')
        else:
            ws.append([item_id, item_code, price, category, category_code, image_path or '无'])
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = '货品列表.xlsx'
    if category_id or search:
        filename = '货品列表_筛选结果.xlsx'
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

@app.route('/api/export/csv', methods=['GET'])
def api_export_csv():
    """导出货品数据为CSV"""
    import csv
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    category_id = request.args.get('category_id', type=int)
    search = request.args.get('search', '')
    
    query = '''
        SELECT i.id, i.item_code, i.price, c.name as category, c.code as category_code, i.image_path
        FROM items i
        JOIN categories c ON i.category_id = c.id
    '''
    params = []
    
    if category_id:
        query += ' WHERE i.category_id = ?'
        params.append(category_id)
    
    if search:
        if category_id:
            query += ' AND i.item_code LIKE ?'
        else:
            query += ' WHERE i.item_code LIKE ?'
        params.append(f'%{search}%')
    
    query += ' ORDER BY c.id, i.id'
    
    cursor.execute(query, params)
    items = cursor.fetchall()
    conn.close()
    
    output = io.BytesIO()
    
    # 使用utf-8-sig编码以确保Excel正确打开
    import codecs
    writer_wrapper = codecs.getwriter('utf-8-sig')
    output_wrapper = writer_wrapper(output)
    
    writer = csv.writer(output_wrapper)
    
    writer.writerow(['ID', '货号', '价格', '分类', '分类编号', '图片路径'])
    writer.writerows(items)
    
    output.seek(0)
    
    return send_file(
        output,
        mimetype='text/csv',
        as_attachment=True,
        download_name='货品列表.csv'
    )

def open_browser():
    """延迟打开浏览器"""
    time.sleep(2)
    webbrowser.open('http://127.0.0.1:3001')

if __name__ == '__main__':
    # 检查模型
    check_model()
    
    # 设置模板和静态文件路径
    app.template_folder = str(BASE_PATH / 'templates')
    app.static_folder = str(BASE_PATH / 'static')
    
    # 预加载模型
    print("正在初始化应用...")
    get_model()
    print("初始化完成！")
    
    # 自动打开浏览器
    if not os.environ.get('FLASK_DEBUG'):
        threading.Thread(target=open_browser, daemon=True).start()
        print("浏览器将自动打开...")
    
    print("=" * 50)
    print("美甲图片匹配器已启动")
    print("访问地址: http://127.0.0.1:3001")
    print("按 Ctrl+C 关闭应用")
    print("=" * 50)
    
    app.run(debug=False, port=3001, host='127.0.0.1')